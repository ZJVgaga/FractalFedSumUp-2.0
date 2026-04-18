import openpyxl
from openpyxl.utils.cell import column_index_from_string, get_column_letter



# 定义一个函数，用于将hyperparam 到 cell address的映射 转换为 cell address到hyperparam的映射
def reverse_hyperparam_to_cell_address_map(hyperparam_to_cell_address_map):
    """
    将 hyperparam_to_cell_address_map 转换成 cell_address_to_hyperparam_map

    输入格式：
        hyperparam_to_cell_address_map: dict
            {
                "config_key_1": "start_cell-end_cell",  # 单个区域
                "config_key_2": ["start_cell1-end_cell1", "start_cell2-end_cell2"],  # 多个区域
                ...
            }

    输出格式：
        cell_address_to_hyperparam_map: dict
            {"A1": "config_key_1", "A2": "config_key_1", ...}
    """
    # 初始化一个空字典，用于存储 cell 到配置项的映射
    
    cell_address_to_hyperparam_map = {}
    
    # 遍历 hyperparam_to_cell_address_map 中的每个配置项及其对应的单元格范围
    for config_key, cell_ranges in hyperparam_to_cell_address_map.items():
        # 如果单元格范围是一个列表（表示多个区域）
        if isinstance(cell_ranges, list):
            # 遍历列表中的每个单元格范围
            for cell_range in cell_ranges:
                # 将单元格范围字符串拆分为起始单元格和结束单元格
                start_cell, end_cell = cell_range.split('-')
                
                # 从起始单元格中提取行号和列号
                row_start = int(start_cell[1:])  # 提取行号（例如 "A1" -> 1）
                col_start = column_index_from_string(start_cell[:1])  # 提取列号（例如 "A1" -> 1）
                
                # 从结束单元格中提取行号和列号
                row_end = int(end_cell[1:])  # 提取行号（例如 "B2" -> 2）
                col_end = column_index_from_string(end_cell[:1])  # 提取列号（例如 "B2" -> 2）

                # 遍历当前单元格范围内的所有行和列
                for row in range(row_start, row_end + 1):  # 遍历行号
                    for col in range(col_start, col_end + 1):  # 遍历列号
                        # 构造单元格地址（例如 "A1"）
                        cell_address = f'{get_column_letter(col)}{row}'
                        
                        # 将单元格地址映射到对应的配置项
                        # 格式：{"A1": "config_key_1", "A2": "config_key_1", ...}
                        cell_address_to_hyperparam_map[cell_address] = config_key
        else:
            # 如果单元格范围不是列表（表示单个区域）
            # 将单元格范围字符串拆分为起始单元格和结束单元格
            start_cell, end_cell = cell_ranges.split('-')
            
            # 从起始单元格中提取行号和列号
            row_start = int(start_cell[1:])  # 提取行号（例如 "A1" -> 1）
            col_start = column_index_from_string(start_cell[:1])  # 提取列号（例如 "A1" -> 1）
            
            # 从结束单元格中提取行号和列号
            row_end = int(end_cell[1:])  # 提取行号（例如 "B2" -> 2）
            col_end = column_index_from_string(end_cell[:1])  # 提取列号（例如 "B2" -> 2）

            # 遍历当前单元格范围内的所有行和列
            for row in range(row_start, row_end + 1):  # 遍历行号
                for col in range(col_start, col_end + 1):  # 遍历列号
                    # 构造单元格地址（例如 "A1"）
                    cell_address = f'{get_column_letter(col)}{row}'
                    
                    # 将单元格地址映射到对应的配置项
                    # 格式：{"A1": "config_key_1", "A2": "config_key_1", ...}
                    cell_address_to_hyperparam_map[cell_address] = config_key
    
    # 返回生成的 cell 到配置项的映射
    # 格式：{"A1": "config_key_1", "A2": "config_key_1", ...}
    return cell_address_to_hyperparam_map